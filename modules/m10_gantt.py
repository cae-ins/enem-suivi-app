"""Diagramme de Gantt des tâches et des points de suivi d'un trimestre.

Code d'origine : generer_gantt.py + Fichier_Gant.xlsx
Décisions : tâches modifiables dans l'application (GA1), points de suivi ajoutés
automatiquement à partir des fréquences des modules (GA2), dépendances vérifiées par
ordre d'exécution (GA3), plus de menu console (GA4).
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

from enem_core import donnees as D
from enem_core.execution import ModuleBase
from enem_core.exports import BORDURE, VERT
from enem_core.parametres import Booleen, DateP, Entier, Fichier, Texte

from ._communs import p_sortie

COLS = ["Timestre_a_debute", "Ordre_exécution", "Intitule_taches", "Date_debut_Tache", "Date_fin_tache"]


def lire_taches(chemin) -> pd.DataFrame:
    if not chemin or not Path(chemin).is_file():
        return pd.DataFrame(columns=COLS)
    df = pd.read_excel(chemin)
    manq = [c for c in COLS if c not in df.columns]
    if manq:
        raise D.ErreurDonnees(f"Colonnes absentes du fichier des tâches : {', '.join(manq)}")
    df["Date_debut_Tache"] = pd.to_datetime(df["Date_debut_Tache"], errors="coerce")
    df["Date_fin_tache"] = pd.to_datetime(df["Date_fin_tache"], errors="coerce")
    df["Timestre_a_debute"] = df["Timestre_a_debute"].astype(str).str.strip()
    return df


def ecrire_taches(chemin, df: pd.DataFrame):
    d = df[COLS].copy()
    d["Date_debut_Tache"] = pd.to_datetime(d["Date_debut_Tache"]).dt.date
    d["Date_fin_tache"] = pd.to_datetime(d["Date_fin_tache"]).dt.date
    Path(chemin).parent.mkdir(parents=True, exist_ok=True)
    d.to_excel(chemin, index=False, sheet_name="Feuil1")


def points_de_suivi(debut: dt.date, nb_semaines: int, trimestre: str) -> pd.DataFrame:
    from modules import charger_modules  # import tardif (évite la dépendance circulaire)
    fin = debut + dt.timedelta(days=7 * nb_semaines - 1)
    lignes = []
    for m in charger_modules():
        if not m.frequence_jours or m.id in ("gantt", "preparation_bases"):
            continue
        jour = debut + dt.timedelta(days=m.frequence_jours - 1)
        while jour <= fin:
            lignes.append({"Timestre_a_debute": trimestre, "Ordre_exécution": 90,
                           "Intitule_taches": f"Point de suivi – {m.nom}",
                           "Date_debut_Tache": pd.Timestamp(jour), "Date_fin_tache": pd.Timestamp(jour)})
            jour += dt.timedelta(days=m.frequence_jours)
    return pd.DataFrame(lignes, columns=COLS)


def verifier_dependances(df: pd.DataFrame) -> pd.DataFrame:
    """Signale les tâches d'ordre n+1 qui commencent avant la fin de toutes les tâches d'ordre n."""
    pb = []
    ordres = sorted(o for o in df["Ordre_exécution"].dropna().unique() if o < 90)
    for a, b in zip(ordres, ordres[1:]):
        fin_a = df.loc[df["Ordre_exécution"] == a, "Date_fin_tache"].max()
        for _, r in df[df["Ordre_exécution"] == b].iterrows():
            if r["Date_debut_Tache"] < fin_a:
                pb.append({"Ordre": int(b), "Tâche": r["Intitule_taches"],
                           "Début": r["Date_debut_Tache"].strftime("%d/%m/%Y"),
                           "Fin des tâches d'ordre précédent": fin_a.strftime("%d/%m/%Y"),
                           "Message": f"Commence avant la fin des tâches d'ordre {int(a)}"})
    return pd.DataFrame(pb, columns=["Ordre", "Tâche", "Début", "Fin des tâches d'ordre précédent", "Message"])


class Gantt(ModuleBase):
    id = "gantt"
    nom = "Diagramme de Gantt du trimestre"
    nom_court = "Diagramme de Gantt"
    description = ("Produit le diagramme de Gantt des tâches d'un trimestre (fichier des tâches modifiable dans "
                   "l'application) en y ajoutant, si souhaité, le calendrier des points de suivi, et vérifie "
                   "l'ordre d'exécution des tâches.")
    equipe = "Coordination"
    frequence_jours = None
    frequence_libelle = "À la demande"
    couleur = "#1F6F8B"
    entrees = "Fichier_Gant.xlsx (tâches) + fréquences des modules"
    sorties = "Diagramme_Gantt_<T>.xlsx (Gantt visuel, graphique, récapitulatif, dépendances)"
    code_origine = ["Code_gant/generer_gantt.py"]
    mots_cles = "gantt planning calendrier tâches points de suivi"
    ordre = 90
    sous_dossier = "Coordination"

    @property
    def parametres(self):
        return [
            Texte("trimestre", "Trimestre à afficher (valeur de Timestre_a_debute)", "T3_2026", groupe="Général",
                  globale="trimestre"),
            Fichier("fichier_taches", "Fichier des tâches (Fichier_Gant.xlsx)", "", groupe="Général",
                    globale="fichier_gantt", types=[("Excel", "*.xlsx")],
                    aide="Modifiable depuis l'écran « Tâches du Gantt »."),
            Booleen("ajouter_points", "Ajouter les points de suivi des modules", True, groupe="Points de suivi"),
            DateP("debut_points", "Début de la collecte (lundi semaine 1)", "", groupe="Points de suivi",
                  globale="lundi_semaine1"),
            Entier("nb_semaines", "Nombre de semaines de référence", 13, groupe="Points de suivi", mini=1, maxi=30),
            p_sortie(),
        ]

    def executer(self, ctx, p):
        tr = str(p["trimestre"]).strip()
        taches = lire_taches(p["fichier_taches"])
        taches = taches[taches["Timestre_a_debute"].str.upper() == tr.upper()]
        if p["ajouter_points"]:
            pts = points_de_suivi(p["debut_points"], p["nb_semaines"], tr)
            taches = pd.concat([taches, pts], ignore_index=True)
            ctx.journal(f"  {len(pts)} point(s) de suivi ajouté(s)")
        taches = taches.dropna(subset=["Date_debut_Tache", "Date_fin_tache"])
        if taches.empty:
            raise D.ErreurDonnees(f"Aucune tâche pour le trimestre {tr}.")
        taches = taches.sort_values(["Ordre_exécution", "Date_debut_Tache"]).reset_index(drop=True)
        ctx.progression(20, "Vérification des dépendances")
        dep = verifier_dependances(taches)
        for _, r in dep.iterrows():
            ctx.alerte(f"Tâche {r['Ordre']} – {r['Tâche'][:50]}", r["Message"], "Moyenne")
        ctx.progression(40, "Construction du diagramme")
        d0, d1 = taches["Date_debut_Tache"].min(), taches["Date_fin_tache"].max()
        jours = pd.date_range(d0, d1, freq="D")
        lignes = []
        for _, t in taches.iterrows():
            l = {"Ordre": int(t["Ordre_exécution"]), "Tâche": t["Intitule_taches"],
                 "Début": t["Date_debut_Tache"].strftime("%d/%m/%Y"), "Fin": t["Date_fin_tache"].strftime("%d/%m/%Y"),
                 "Durée (j)": (t["Date_fin_tache"] - t["Date_debut_Tache"]).days + 1}
            for j in jours:
                l[j.strftime("%d/%m/%y")] = "■" if t["Date_debut_Tache"] <= j <= t["Date_fin_tache"] else ""
            lignes.append(l)
        visuel = pd.DataFrame(lignes)
        ouvrables = int((jours.weekday < 5).sum())
        recap = pd.DataFrame([
            ("Trimestre", tr), ("Nombre de tâches", len(taches)),
            ("Date de début", d0.strftime("%d/%m/%Y")), ("Date de fin", d1.strftime("%d/%m/%Y")),
            ("Durée totale (jours calendaires)", len(jours)), ("Durée totale (jours ouvrables)", ouvrables),
            ("Durée moyenne par tâche (jours)", round(visuel["Durée (j)"].mean(), 1)),
        ], columns=["Catégorie", "Valeur"])
        liste = taches.assign(Debut=taches["Date_debut_Tache"].dt.date, Fin=taches["Date_fin_tache"].dt.date)[
            ["Ordre_exécution", "Intitule_taches", "Debut", "Fin"]]
        chemin = ctx.chemin(f"Diagramme_Gantt_{tr}.xlsx")
        with pd.ExcelWriter(chemin, engine="openpyxl") as w:
            visuel.to_excel(w, sheet_name="Gantt Visuel", index=False, startrow=1)
            recap.to_excel(w, sheet_name="Récapitulatif", index=False)
            liste.to_excel(w, sheet_name="Liste des tâches", index=False)
            dep.to_excel(w, sheet_name="Dépendances", index=False)
        ctx.progression(70, "Mise en forme Excel")
        self._mise_en_forme(chemin, jours, tr, taches)
        ctx.fichier_produit(chemin)
        ctx.journal(f"  ✓ Fichier créé : {chemin.name}")
        ctx.resultat.tables.update({"Liste des tâches": liste, "Récapitulatif": recap, "Dépendances": dep})
        ctx.graphique(f"Gantt – {tr}", "Liste des tâches", "Debut", ["Fin"], "gantt", etiquette="Intitule_taches")
        ctx.resultat.resume = f"{len(taches)} tâche(s), {len(dep)} problème(s) d'ordre – {tr}"

    @staticmethod
    def _mise_en_forme(chemin, jours, tr, taches):
        wb = load_workbook(chemin)
        ws = wb["Gantt Visuel"]
        vert = PatternFill("solid", start_color=VERT, end_color=VERT)
        orange = PatternFill("solid", start_color="F39422", end_color="F39422")
        gris = PatternFill("solid", start_color="D9D9D9", end_color="D9D9D9")
        ws["A1"] = f"DIAGRAMME DE GANTT – {tr}"
        ws["A1"].font = Font(bold=True, size=14, color=VERT)
        for c in ws[2]:
            c.fill, c.font = vert, Font(bold=True, color="FFFFFF", size=9)
            c.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90 if c.column > 5 else 0)
            c.border = BORDURE
        for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
            point = str(row[1].value or "").startswith("Point de suivi")
            for c in row:
                c.border = BORDURE
                if c.column > 5:
                    j = jours[c.column - 6]
                    if j.weekday() >= 5:
                        c.fill = gris
                    if c.value == "■":
                        c.fill = orange if point else vert
                        c.font = Font(color="F39422" if point else VERT)
        ws.column_dimensions["A"].width = 7
        ws.column_dimensions["B"].width = 55
        for col in "CDE":
            ws.column_dimensions[col].width = 11
        from openpyxl.utils import get_column_letter
        for k in range(6, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(k)].width = 3.2
        ws.row_dimensions[2].height = 48
        ws.freeze_panes = "F3"
        ws2 = wb["Récapitulatif"]
        ws2.column_dimensions["A"].width = 38
        ws2.column_dimensions["B"].width = 25
        for c in ws2[1]:
            c.fill, c.font = vert, Font(bold=True, color="FFFFFF")
        # graphique en barres empilées (première série invisible)
        wg = wb.create_sheet("Graphique Gantt")
        wg.append(["Tâche", "Jours avant début", "Durée (jours)"])
        d0 = taches["Date_debut_Tache"].min()
        for _, t in taches.iterrows():
            wg.append([f"#{int(t['Ordre_exécution'])} {str(t['Intitule_taches'])[:40]}",
                       (t["Date_debut_Tache"] - d0).days, (t["Date_fin_tache"] - t["Date_debut_Tache"]).days + 1])
        ch = BarChart()
        ch.type, ch.grouping, ch.overlap = "bar", "stacked", 100
        ch.title = f"Diagramme de Gantt – {tr}"
        ch.height, ch.width = max(8, 0.5 * len(taches)), 26
        n = len(taches) + 1
        ch.add_data(Reference(wg, min_col=2, min_row=1, max_col=3, max_row=n), titles_from_data=True)
        ch.set_categories(Reference(wg, min_col=1, min_row=2, max_row=n))
        ch.series[0].graphicalProperties.noFill = True
        ch.series[0].graphicalProperties.line.noFill = True
        ch.series[1].graphicalProperties.solidFill = VERT
        ch.y_axis.scaling.orientation = "minMax"
        ch.x_axis.scaling.orientation = "maxMin"
        ch.legend = None
        wg.add_chart(ch, "E2")
        wb.save(chemin)


MODULE = Gantt()
