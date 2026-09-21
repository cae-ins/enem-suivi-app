"""Construction des variables du bulletin et moteurs de tabulation.

Portage fidèle des do-files de la chaîne « Simulation indicateur » :

* ``1_1_Var_objectives_lower.do``        → :func:`variables_objectives`
* ``1_2_Indicateur_Bulletin_To_Run.do``  → :func:`variables_bulletin`
* ``Revision_CISE_12112024.do``          → :func:`variables_cise`
* ``programme_master_simple.do``         → :func:`indicateur_simple`
* ``programme_master_simple_ANNUEL.do``  → :func:`indicateur_simple` (periode="annee")
* ``programme_repartition_horiz_FINAL_total.do``    → :func:`repartition_horiz`
* ``programme_repartition_horiz_CORRIGE_ANNUEL.do`` → :func:`repartition_horiz`
  (avec ``lignes_corrigees=True`` : correction des bugs 1 et 2 de décalage de lignes)

Sémantique Stata respectée
--------------------------
Dans Stata, une valeur manquante vaut *plus l'infini* dans les comparaisons :
``age >= 16`` est vrai si ``age`` est manquant. Les fonctions :func:`ge`, :func:`gt`,
:func:`le`, :func:`lt` reproduisent ce comportement, tandis que :func:`inrange` et
:func:`inlist` excluent les manquants, exactement comme les fonctions Stata
homonymes. Les résultats sont donc identiques à ceux des do-files, y compris pour
les observations à âge manquant.
"""

from __future__ import annotations

import re
import shutil
from pathlib import Path

import numpy as np
import pandas as pd

# ---------------------------------------------------------------------------
# Sémantique Stata
# ---------------------------------------------------------------------------


def col(df: pd.DataFrame, nom: str) -> pd.Series:
    """Colonne numérique ``nom`` ; série de manquants si la variable est absente."""
    if nom not in df.columns:
        return pd.Series(np.nan, index=df.index, dtype=float)
    return pd.to_numeric(df[nom], errors="coerce")


def texte(df: pd.DataFrame, nom: str) -> pd.Series:
    if nom not in df.columns:
        return pd.Series("", index=df.index, dtype=object)
    return df[nom].map(lambda v: "" if pd.isna(v) else str(v))


def ge(s: pd.Series, v) -> pd.Series:
    """``s >= v`` au sens Stata (manquant = +∞ donc vrai)."""
    return s.isna() | (s >= v)


def gt(s: pd.Series, v) -> pd.Series:
    return s.isna() | (s > v)


def le(s: pd.Series, v) -> pd.Series:
    return (~s.isna()) & (s <= v)


def lt(s: pd.Series, v) -> pd.Series:
    return (~s.isna()) & (s < v)


def inrange(s: pd.Series, a, b) -> pd.Series:
    """``inrange()`` de Stata : faux pour les manquants."""
    return (~s.isna()) & (s >= a) & (s <= b)


def inlist(s: pd.Series, *valeurs) -> pd.Series:
    """``inlist()`` de Stata : faux pour les manquants."""
    return s.isin(list(valeurs)) & (~s.isna())


def eq(s: pd.Series, v) -> pd.Series:
    return (~s.isna()) & (s == v)


def ne(s: pd.Series, v) -> pd.Series:
    """``s != v`` au sens Stata : vrai lorsque s est manquant."""
    return s.isna() | (s != v)


def present(s: pd.Series) -> pd.Series:
    """``!missing(s)``."""
    return ~s.isna()


def vide(n: int, index) -> pd.Series:
    return pd.Series(np.nan, index=index, dtype=float)


def poser(serie: pd.Series, masque: pd.Series, valeur) -> pd.Series:
    """Équivalent de ``replace var = valeur if masque``."""
    serie = serie.copy()
    serie.loc[masque.fillna(False)] = valeur
    return serie


def normaliser_noms(df: pd.DataFrame) -> pd.DataFrame:
    """Minuscules + ``__`` → ``_`` (équivalent Stata de ``normalize_column_names``)."""
    nouveaux, conflits = {}, []
    existants = set(df.columns)
    for c in df.columns:
        n = c.lower()
        while "__" in n:
            n = n.replace("__", "_")
        if n != c:
            if n in existants or n in nouveaux.values():
                conflits.append(f"{c}->{n}")
            else:
                nouveaux[c] = n
    if nouveaux:
        df = df.rename(columns=nouveaux)
    df.attrs["conflits_noms"] = conflits
    return df


# ---------------------------------------------------------------------------
# 1_1_Var_objectives_lower.do
# ---------------------------------------------------------------------------

ABIDJAN = [1010100211, 1010100212, 1010100213, 1010100214, 1010100215,
           1010100216, 1010100217, 1010100218, 1010100219, 1010100220]


def _recode(s: pd.Series, regles: list) -> pd.Series:
    """``recode`` Stata : valeurs non citées conservées, manquants conservés."""
    r = s.copy().astype(float)
    for sources, cible in regles:
        r = poser(r, s.isin(list(sources)), cible)
    return r


def variables_objectives(df: pd.DataFrame) -> pd.DataFrame:
    """Portage de ``1_1_Var_objectives_lower.do``."""
    i = df.index
    ageannee, m4 = col(df, "ageannee"), col(df, "m4confirm")

    age = pd.Series(np.nan, index=i, dtype=float)
    age = poser(age, ge(ageannee, 13), m4)          # age = m4confirm if ageannee >= 13
    age.loc[lt(ageannee, 13)] = ageannee[lt(ageannee, 13)]
    df["age"] = age

    df["grp_age"] = np.where(
        present(age),
        1 * lt(age, 15) + 2 * (ge(age, 15) & le(age, 24)) + 3 * (gt(age, 24) & le(age, 35))
        + 4 * (gt(age, 35) & le(age, 64)) + 5 * gt(age, 64), np.nan)

    bornes5 = [(None, 4), (4, 9), (9, 14), (14, 19), (19, 24), (24, 29), (29, 34), (34, 39),
               (39, 44), (44, 49), (49, 54), (54, 59), (59, 64)]
    g5 = pd.Series(0.0, index=i)
    for k, (bas, haut) in enumerate(bornes5, start=1):
        m = le(age, haut) if bas is None else (gt(age, bas) & le(age, haut))
        g5 += k * m
    g5 += 14 * gt(age, 64)
    df["grpe_age5"] = np.where(present(age), g5, np.nan)

    df["grp_age3"] = np.where(
        present(age),
        1 * lt(age, 15) + 2 * (ge(age, 15) & le(age, 24)) + 3 * (gt(age, 24) & le(age, 64))
        + 4 * gt(age, 64), np.nan)

    df["groupe_age4"] = (1 * (ge(age, 16) & le(age, 24)) + 2 * (ge(age, 25) & le(age, 35))
                         + 3 * (ge(age, 36) & le(age, 64)) + 4 * ge(age, 65)).astype(float)
    df["groupe_age5"] = (1 * (ge(age, 16) & le(age, 35)) + 2 * (ge(age, 36) & le(age, 64))
                         + 3 * ge(age, 65)).astype(float)
    df["groupe_age6"] = (1 * (ge(age, 15) & le(age, 19)) + 2 * (ge(age, 20) & le(age, 24))
                         + 3 * (ge(age, 25) & le(age, 29)) + 4 * (ge(age, 30) & le(age, 35))
                         + 5 * (ge(age, 36) & le(age, 40))).astype(float)
    df["groupe_age7"] = (1 * (ge(age, 15) & le(age, 19)) + 2 * (ge(age, 20) & le(age, 24))
                         + 3 * (ge(age, 25) & le(age, 29)) + 4 * (ge(age, 30) & le(age, 34))
                         + 5 * (ge(age, 35) & le(age, 40))).astype(float)
    df["jeune15_24"] = (ge(age, 15) & le(age, 24)).astype(float)
    df["jeune15_35"] = (ge(age, 15) & le(age, 35)).astype(float)
    df["jeune15_40"] = (ge(age, 15) & le(age, 40)).astype(float)

    df["sexe"] = col(df, "m5")

    ef1, ef3, ef6, ef7, ef10, ef11 = (col(df, c) for c in ("ef1", "ef3", "ef6", "ef7", "ef10", "ef11"))
    ef4, ef8, ef12 = col(df, "ef4"), col(df, "ef8"), col(df, "ef12")
    a3 = ge(age, 3)

    sco = vide(len(df), i)
    sco = poser(sco, a3 & eq(ef1, 2), 1)
    sco = poser(sco, a3 & eq(ef1, 1), 0)
    df["scolarise"] = sco

    niv = vide(len(df), i)
    niv = poser(niv, eq(ef1, 2) & a3, -99)
    m = eq(ef6, 1) & eq(ef1, 1) & a3
    niv.loc[m] = ef7[m]
    m = eq(ef10, 1) & ne(ef6, 1) & eq(ef1, 1) & a3
    niv.loc[m] = ef11[m]
    m = ne(ef6, 1) & ne(ef10, 1) & eq(ef1, 1) & a3
    niv.loc[m] = ef3[m]
    df["niveau_instruction"] = niv

    df["niv_inst_ag1"] = _recode(niv, [([1, -99], 1), ([2], 2), ([3, 4], 3), ([5, 6], 4),
                                       ([7, 8], 5), ([9, 10], 6), ([11, 12], 7)])
    df["niv_inst_ag2"] = _recode(niv, [([1, -99], 1), ([2], 2), ([3, 4, 5, 6], 3),
                                       ([7, 8, 9, 10, 11, 12], 4)])
    df["niv_inst_ag3"] = _recode(col(df, "niv_inst_ag1"), [([1], 1), ([2], 2), ([3], 3), ([4], 4),
                                                           ([5, 6, 7], 5)])

    cl = vide(len(df), i)
    cl = poser(cl, eq(ef1, 2) & a3, -99)
    m = eq(ef6, 1) & eq(ef1, 1) & a3
    cl.loc[m] = ef8[m]
    m = eq(ef10, 1) & ne(ef6, 1) & eq(ef1, 1) & a3
    cl.loc[m] = ef12[m]
    m = ne(ef10, 1) & ne(ef6, 1) & eq(ef1, 1) & a3
    cl.loc[m] = ef4[m]
    df["classe_atteint"] = cl
    df["nbr_annee_etude"] = _recode(cl, [([-99, 1, 2, 3, 4], 0)] +
                                    [([k], k - 4) for k in range(5, 23)])

    if "ef4_1" in df.columns:
        df["haut_diplome"] = col(df, "ef4_1").where(ge(age, 3))

    hh6, hh4 = col(df, "hh6"), col(df, "hh4")
    df["milieu_residence"] = hh6
    mr = vide(len(df), i)
    mr = poser(mr, eq(hh6, 1), 2)
    mr = poser(mr, eq(hh6, 1) & inlist(hh4, *ABIDJAN), 1)
    mr = poser(mr, eq(hh6, 2), 3)
    df["milieu_resid2"] = mr

    for cible, source in (("region", "hh2"), ("district", "hh1"), ("type_logement", "l1"),
                          ("nature_murs", "l3"), ("statut_occupation", "l5"),
                          ("mode_eclairage", "l4"), ("cat_profEP", "ep13"), ("cat_profES", "es13")):
        if source in df.columns:
            df[cible] = col(df, source)
    return df


# ---------------------------------------------------------------------------
# 1_2_Indicateur_Bulletin_To_Run.do
# ---------------------------------------------------------------------------


def variables_bulletin(df: pd.DataFrame) -> pd.DataFrame:
    """Portage de ``1_2_Indicateur_Bulletin_To_Run.do``."""
    i = df.index
    c = lambda n: col(df, n)  # noqa: E731
    age = c("age")
    PAT = (ge(age, 16) & present(age)).astype(float)
    df["PAT"] = PAT
    pat = eq(PAT, 1)

    se = {k: c(f"se{k}") for k in ("1", "2", "3", "4", "5", "7", "8", "9", "9a", "9b", "10", "11")}
    emp_p = pd.Series(np.nan, index=i, dtype=float)
    emp_p = poser(emp_p, pat, 0)
    for m in (eq(se["1"], 1),
              eq(se["1"], 2) & inrange(se["2"], 1, 9),
              eq(se["2"], 10) & inrange(se["3"], 1, 2),
              inrange(se["3"], 3, 5) & eq(se["4"], 1),
              eq(se["4"], 2) & eq(se["5"], 1) & inlist(se["7"], 1, 2)):
        emp_p = poser(emp_p, m & pat, 1)
    df["emp_present"] = emp_p

    absent_base = eq(se["5"], 2) | inlist(se["7"], 3, 4)
    se9_autres = inrange(se["9"], 5, 9) | inlist(se["9"], 11, 12, 15, 16, 17)
    emp_a = pd.Series(np.nan, index=i, dtype=float)
    emp_a = poser(emp_a, pat, 0)
    for m in (absent_base & eq(se["8"], 1) & inrange(se["9"], 1, 4),
              absent_base & eq(se["8"], 1) & eq(se["9"], 10) & eq(se["9a"], 1),
              absent_base & eq(se["8"], 1) & eq(se["9"], 14) & eq(se["9b"], 1),
              absent_base & eq(se["8"], 1) & se9_autres & eq(se["10"], 1),
              absent_base & eq(se["8"], 1) & se9_autres & eq(se["10"], 2) & eq(se["11"], 1)):
        emp_a = poser(emp_a, m & pat, 1)
    df["emp_absent"] = emp_a

    pop_emp = pd.Series(np.nan, index=i, dtype=float)
    pop_emp = poser(pop_emp, pat, 0)
    pop_emp = poser(pop_emp, eq(emp_p, 1) & pat, 1)
    pop_emp = poser(pop_emp, eq(emp_a, 1) & pat, 2)
    df["pop_emp"] = pop_emp
    df["pop_emp_dich"] = _recode(pop_emp, [([1, 2], 1)])
    ped = c("pop_emp_dich")

    ep11 = c("ep11")
    df["secteur_institutionnel"] = np.where(
        (eq(ped, 1) & present(ep11)),
        1 * inlist(ep11, 1, 2) + 2 * inlist(ep11, 3, 4) + 3 * inlist(ep11, 7, 8)
        + 4 * eq(ep11, 9) + 5 * eq(ep11, 6), np.nan)
    si2 = pd.Series(np.nan, index=i, dtype=float)
    si2 = poser(si2, pat & inlist(ped, 1, 2), 2)
    si2 = poser(si2, inlist(ep11, 1, 2) & pat & inlist(ped, 1, 2), 1)
    si2 = poser(si2, eq(ep11, 8) & pat & inlist(ped, 1, 2), 3)
    df["secteur_institionnel2"] = si2
    df["secteur_institutionnel3"] = _recode(ep11, [([1, 2], 1), ([3], 2), ([4], 3), ([5], 4),
                                                   ([6, 7], 5), ([8], 6)])

    srh = {k: c(f"srh{k}") for k in ("1", "2", "2a", "6", "7", "8", "9", "11")}
    aucun_emp = (eq(ped, 0) & ((eq(srh["1"], 1) | eq(srh["2"], 1) | eq(srh["2a"], 1)) & eq(srh["11"], 1))
                 & pat).astype(float)
    df["aucun_emp"] = aucun_emp
    futures = (eq(srh["1"], 2) & eq(srh["2"], 2) & eq(srh["2a"], 2) & eq(srh["7"], 18)
               & inlist(srh["8"], 1, 2) & eq(srh["9"], 1) & ~inlist(pop_emp, 1, 2) & pat).astype(float)
    df["futures_staters"] = futures

    pop_cho = pd.Series(np.nan, index=i, dtype=float)
    pop_cho = poser(pop_cho, eq(aucun_emp, 1) & eq(ped, 0) & pat, 1)
    pop_cho = poser(pop_cho, eq(futures, 1) & eq(ped, 0) & pat, 2)
    df["pop_chomage"] = pop_cho
    df["pop_chomage_dich"] = _recode(pop_cho, [([1, 2], 1)])
    pcd = c("pop_chomage_dich")

    statut = pd.Series(np.nan, index=i, dtype=float)
    statut = poser(statut, pat, 3)
    statut = poser(statut, eq(ped, 1) & pat, 1)
    statut = poser(statut, eq(pcd, 1) & pat, 2)
    df["statut_MO"] = statut

    MO = pd.Series(np.nan, index=i, dtype=float)
    MO = poser(MO, pat & eq(statut, 1), 1)
    MO = poser(MO, pat & eq(statut, 2), 2)
    df["MO"] = MO
    df["MO_dich"] = _recode(MO, [([1, 2], 1)])

    non_dispo = (pat & ~inlist(pop_emp, 1, 2) & ~inlist(pop_cho, 1, 2)
                 & (eq(srh["1"], 1) | eq(srh["2"], 1) | eq(srh["2a"], 1)) & eq(srh["11"], 2)).astype(float)
    df["Non_dispo"] = non_dispo
    aucune_rech = (pat & ~inlist(pop_emp, 1, 2) & ~inlist(pop_cho, 1, 2)
                   & (eq(srh["1"], 2) & eq(srh["2"], 2) & eq(srh["2a"], 2))
                   & eq(srh["6"], 1) & eq(srh["11"], 1)).astype(float)
    df["aucune_rech"] = aucune_rech

    MOPOT = pd.Series(np.nan, index=i, dtype=float)
    MOPOT = poser(MOPOT, pat & eq(non_dispo, 1), 1)
    MOPOT = poser(MOPOT, pat & eq(aucune_rech, 1), 2)
    df["MOPOT"] = MOPOT
    df["MOPOT_dich"] = _recode(MOPOT, [([1, 2], 1)])
    mb = _recode(MOPOT, [([1, 2], 1)])
    mb = poser(mb, ne(mb, 1) & eq(statut, 3), 0)
    df["MOPOT_bis"] = mb
    df["workers_decu"] = (eq(aucune_rech, 1) & eq(srh["7"], 1)).astype(float)
    df["Non_demand"] = (pat & ~inlist(pop_emp, 1, 2) & eq(srh["1"], 2) & eq(srh["2"], 2)
                        & eq(srh["2a"], 2) & eq(srh["11"], 2) & eq(srh["6"], 1)).astype(float)

    MOE = pd.Series(np.nan, index=i, dtype=float)
    MOE = poser(MOE, pat & inlist(MO, 1, 2), 1)
    MOE = poser(MOE, pat & inlist(MOPOT, 1, 2), 2)
    df["MOE"] = MOE
    df["MOE_dich"] = _recode(MOE, [([1, 2], 1)])

    df["hor_eff"] = c("nb_heure_travail_total")
    hor = c("hor_eff")
    sous = pd.Series(np.nan, index=i, dtype=float)
    sous = poser(sous, pat & inlist(pop_emp, 1, 2), 0)
    sous = poser(sous, pat & inlist(pop_emp, 1, 2) & lt(hor, 40) & eq(c("wki4"), 1) & eq(c("wki5"), 1), 1)
    df["sous_emp"] = sous

    su1 = pd.Series(np.nan, index=i, dtype=float)
    su1 = poser(su1, inlist(MO, 1, 2) & pat, 0)
    su1 = poser(su1, eq(pcd, 1) & pat & inlist(MO, 1, 2), 1)
    df["SU1"] = su1
    su2 = pd.Series(np.nan, index=i, dtype=float)
    su2 = poser(su2, pat & inlist(MO, 1, 2), 0)
    su2 = poser(su2, (eq(pcd, 1) | eq(sous, 1)) & pat & inlist(MO, 1, 2), 1)
    df["SU2"] = su2
    su3 = pd.Series(np.nan, index=i, dtype=float)
    su3 = poser(su3, pat & inlist(MOE, 1, 2), 0)
    su3 = poser(su3, (eq(pcd, 1) | inlist(MOPOT, 1, 2)) & pat & inlist(MOE, 1, 2), 1)
    df["SU3"] = su3
    su4 = pd.Series(np.nan, index=i, dtype=float)
    su4 = poser(su4, pat & inlist(MOE, 1, 2), 0)
    su4 = poser(su4, (eq(pcd, 1) | inlist(MOPOT, 1, 2) | eq(sous, 1)) & pat & inlist(MOE, 1, 2), 1)
    df["SU4"] = su4

    # --- situation dans l'emploi (version minuscules du do-file) ---
    ep3, ep4, ep5, ep6a, ep10c = c("ep3"), c("ep4"), c("ep5"), c("ep6a"), c("ep10c")
    ep14b, ep28, ep29, ep30, ep30b, ep37 = (c("ep14b"), c("ep28"), c("ep29"), c("ep30"),
                                            c("ep30b"), c("ep37"))
    emploi = inlist(pop_emp, 1, 2) & pat

    duree = pd.Series(np.nan, index=i, dtype=float)
    duree.loc[eq(ep30b, 3)] = ep30[eq(ep30b, 3)]
    duree.loc[eq(ep30b, 2)] = (ep30 * 365)[eq(ep30b, 2)]
    duree.loc[eq(ep30b, 1)] = (ep30 * 30)[eq(ep30b, 1)]
    df["duree_contrat"] = duree

    sit = pd.Series(np.nan, index=i, dtype=float)
    sit = poser(sit, emploi, 99)
    sit = poser(sit, eq(ep3, 2) & (inlist(ep11, 1, 2, 5) | (inlist(ep11, 3, 4, 6, 7) & eq(ep14b, 1))) & emploi, 11)
    sit = poser(sit, eq(ep3, 2) & (inlist(ep11, 3, 4, 6, 7, 8) & inlist(ep14b, 2, 9998)) & emploi, 12)
    sit = poser(sit, eq(ep3, 3) & (inlist(ep11, 3, 4, 6, 7) & eq(ep14b, 1)) & emploi, 21)
    sit = poser(sit, eq(ep3, 3) & (eq(ep11, 8) | (inlist(ep11, 3, 4, 6, 7) & inlist(ep14b, 2, 9998))) & emploi, 22)
    sit = poser(sit, eq(ep3, 1) & ne(ep6a, 1) & eq(ep37, 2) & ne(ep10c, 1) & emploi, 3)
    sit = poser(sit, (eq(ep3, 1) & inlist(ep29, 3, 4)) | (eq(ep3, 1) & eq(ep37, 1) & emploi), 41)
    sit = poser(sit, eq(ep3, 1) & ((inlist(ep29, 1, 2) & ge(duree, 365)) | (eq(ep28, 9998) & eq(ep37, 1))) & emploi, 42)
    sit = poser(sit, eq(ep3, 1) & ((inlist(ep29, 1, 2) & lt(duree, 365)) | eq(ep28, 9998)) & emploi, 43)
    sit = poser(sit, inlist(ep3, 8, 9), 44)
    sit = poser(sit, inlist(ep3, 5, 6), 5)
    sit = poser(sit, eq(ep3, 4), 6)
    sit = poser(sit, eq(ep3, 10), 7)
    df["sit_empEP"] = sit

    s2 = pd.Series(np.nan, index=i, dtype=float)
    for valeurs, cible in (((11, 12), 1), ((21, 22), 2), ((3,), 3), ((41, 42, 43, 44), 4),
                           ((5,), 5), ((6,), 6), ((7,), 7)):
        s2 = poser(s2, inlist(sit, *valeurs), cible)
    df["sit_empEP2"] = s2
    s3 = pd.Series(np.nan, index=i, dtype=float)
    for valeurs, cible in (((11, 12, 21, 22), 1), ((3, 41, 42, 43, 44, 5), 2), ((6,), 3), ((7,), 4)):
        s3 = poser(s3, inlist(sit, *valeurs), cible)
    df["sit_empEP3"] = s3

    df["emp_vul"] = poser(poser(pd.Series(np.nan, index=i, dtype=float), emploi, 0),
                          inlist(ep3, 3, 5) & emploi, 1)
    df["emp_prec"] = poser(poser(pd.Series(np.nan, index=i, dtype=float), emploi, 0),
                           inlist(ep29, 1, 2, 4) & emploi, 1)

    ageannee, ef1, ef6 = c("ageannee"), c("ef1"), c("ef6")
    fp1, fp6 = c("fp1"), c("fp6")
    ne_ = poser(pd.Series(np.nan, index=i, dtype=float), ge(ageannee, 3), 0)
    ne_ = poser(ne_, (ne(ef6, 1) | eq(ef1, 2)) & ge(ageannee, 3), 1)
    df["no_education"] = ne_
    nf = poser(pd.Series(np.nan, index=i, dtype=float), ge(ageannee, 6), 0)
    nf = poser(nf, (eq(fp1, 2) | (eq(fp1, 1) & ne(fp6, 2))) & ge(ageannee, 6), 1)
    df["no_formation"] = nf
    df["Formation_pro"] = poser(pd.Series(np.nan, index=i, dtype=float),
                                (eq(fp1, 1) & eq(fp6, 2)) | eq(fp1, 2), 1)

    neets = (~inlist(pop_emp, 1, 2) & eq(ne_, 1) & eq(nf, 1)).astype(float)
    df["NEETs"] = neets
    for suffixe, jeune in (("15_24", "jeune15_24"), ("15_35", "jeune15_35"), ("15_40", "jeune15_40")):
        j = eq(c(jeune), 1)
        df[f"NEET{suffixe}"] = neets.where(j)
        bis = poser(pd.Series(np.nan, index=i, dtype=float), eq(neets.where(j), 1), 1)
        bis = poser(bis, j & ne(bis, 1), 0)
        df[f"NEET{suffixe}_bis"] = bis

    pl1, pl2 = c("pl1"), c("pl2")
    plu = poser(pd.Series(np.nan, index=i, dtype=float), eq(ped, 1) & pat, 0)
    plu = poser(plu, (eq(pl1, 1) | (eq(pl1, 2) & eq(pl2, 1))) & pat & eq(ped, 1), 1)
    df["pluriactivite"] = plu
    return df


# ---------------------------------------------------------------------------
# Revision_CISE_12112024.do
# ---------------------------------------------------------------------------


def variables_cise(df: pd.DataFrame) -> pd.DataFrame:
    """Portage de ``Revision_CISE_12112024.do``."""
    i = df.index
    c = lambda n: col(df, n)  # noqa: E731
    pat, ped = eq(c("PAT"), 1), eq(c("pop_emp_dich"), 1)
    base = pat & ped
    ep3, ep4, ep5, ep6_1, ep10c, ep11 = (c("ep3"), c("ep4"), c("ep5"), c("ep6_1"), c("ep10c"), c("ep11"))
    ep14b, ep20, ep22, ep26a, ep27a, ep26b, ep27b = (c("ep14b"), c("ep20"), c("ep22"), c("ep26a"),
                                                     c("ep27a"), c("ep26b"), c("ep27b"))
    ep29, ep30c, ep32, ep33, ep35, ep37, ep38, ep39 = (c("ep29"), c("ep30c"), c("ep32"), c("ep33"),
                                                       c("ep35"), c("ep37"), c("ep38"), c("ep39"))
    ep31 = {k: c(f"ep31_{k}") for k in (1, 2, 3, 4, 5)}

    n = pd.Series(np.nan, index=i, dtype=float)
    n = poser(n, inlist(ep3, 2, 3, 4, 10) & eq(ep5, 1) & base, 1)
    n = poser(n, eq(ep3, 5) & inlist(ep4, 1, 2) & eq(ep5, 1) & base, 1)
    n = poser(n, inlist(ep3, 2, 3, 4, 10) & ~inlist(ep5, 1) & base, 2)
    n = poser(n, eq(ep3, 5) & inlist(ep4, 1, 2) & ~inlist(ep5, 1) & base, 2)
    n = poser(n, inlist(ep3, 1, 8, 9, 4, 10) & ne(ep6_1, 1) & ne(ep37, 1) & ne(ep10c, 1) & base, 3)
    n = poser(n, eq(n, 2) & ((eq(ep26a, 2) & inlist(ep27a, 1, 2))
                             | (eq(ep26b, 2) & inlist(ep27b, 1, 2))) & base, 3)
    salarie = eq(ep6_1, 1) | eq(ep37, 1) | eq(ep10c, 1)
    n = poser(n, (inlist(ep3, 4, 5, 6, 10) & ~inlist(ep4, 1, 2)) & salarie & base, 4)
    n = poser(n, inlist(ep3, 1, 4, 10) & salarie & base, 4)
    n = poser(n, inlist(ep3, 8, 9, 10) & base, 4)
    n = poser(n, inlist(ep3, 4, 5, 10) & ~inlist(ep4, 1, 2) & ne(ep6_1, 1) & base, 5)
    n = poser(n, inlist(ep3, 4, 6, 10) & ne(ep6_1, 1) & base, 5)
    df["CISE_18_new"] = n

    hors89 = ~inlist(ep3, 8, 9)
    v = pd.Series(np.nan, index=i, dtype=float)
    v = poser(v, eq(n, 1) & eq(ep14b, 1) & base, 1)
    v = poser(v, base & eq(n, 1) & ne(ep14b, 1), 2)
    v = poser(v, eq(n, 2) & eq(ep14b, 1) & base, 3)
    v = poser(v, ne(ep14b, 1) & base & eq(n, 2), 4)
    v = poser(v, eq(n, 5) & base, 5)
    v = poser(v, eq(n, 3) & base, 6)
    v = poser(v, eq(n, 4) & ~inlist(ep29, 3, 4) & eq(ep31[2], 1) & base & hors89, 7)
    v = poser(v, eq(n, 4) & inlist(ep3, 8, 9) & base, 7)
    v = poser(v, eq(n, 4) & base & hors89, 10)
    perm = inlist(ep29, 3, 4)
    v = poser(v, eq(n, 4) & (perm & eq(ep33, 1)) & base & hors89, 8)
    v = poser(v, eq(n, 4) & (perm & ne(ep33, 1) & eq(ep35, 1)) & base & hors89, 8)
    e31345 = eq(ep31[3], 1) | eq(ep31[4], 1) | eq(ep31[5], 1)
    v = poser(v, eq(n, 4) & (~perm & e31345 & eq(ep32, 1) & eq(ep33, 1)) & base & hors89, 8)
    v = poser(v, eq(n, 4) & (~perm & e31345 & eq(ep32, 1) & ne(ep33, 1) & eq(ep35, 1)) & base & hors89, 8)
    horsc = ~inlist(ep30c, 1, 2, 3, 8)
    v = poser(v, eq(n, 4) & (~perm & horsc & eq(ep31[1], 1) & ne(ep33, 1) & ne(ep35, 2)) & base & hors89, 9)
    v = poser(v, eq(n, 4) & (~perm & horsc & eq(ep31[1], 1) & ne(ep33, 2)) & base & hors89, 9)
    v = poser(v, eq(n, 4) & ~perm & horsc & e31345 & ne(ep32, 1) & ne(ep33, 2) & base & hors89, 9)
    v = poser(v, eq(n, 4) & (~perm & horsc & e31345 & ne(ep32, 1) & ne(ep33, 1) & ne(ep35, 2)) & base & hors89, 9)
    df["CISE_18_niv2"] = v

    aut = pd.Series(np.nan, index=i, dtype=float)
    aut = poser(aut, inlist(c("pop_emp"), 1, 2) & pat, 3)
    aut = poser(aut, inlist(n, 1, 2), 1)
    aut = poser(aut, inlist(n, 3, 4, 5), 2)
    df["sit_empEP_Autorite"] = aut

    legal = inlist(ep14b, 1) | eq(ep20, 1) | eq(ep22, 1)
    inf = pd.Series(np.nan, index=i, dtype=float)
    inf = poser(inf, base & present(ep3) & present(ep11), 2)
    inf = poser(inf, eq(v, 5) & inlist(ep11, 8) & base, 3)
    inf = poser(inf, eq(n, 4) & inlist(ep11, 8) & base, 3)
    inf = poser(inf, eq(n, 3) & inlist(ep11, 8) & base, 3)
    inf = poser(inf, inlist(n, 4) & inlist(ep11, 1, 2, 5) & base, 1)
    inf = poser(inf, inlist(n, 4, 5) & inlist(ep11, 4, 3, 6, 7) & legal & base, 1)
    inf = poser(inf, eq(n, 3) & eq(ep3, 1) & inlist(ep11, 1, 2, 5) & base, 1)
    inf = poser(inf, eq(n, 3) & inlist(ep11, 4, 3, 6, 7) & legal & base, 1)
    inf = poser(inf, inlist(n, 1, 2) & inlist(ep11, 4, 3, 6, 7) & legal & base, 1)
    df["CISE_18_informel"] = inf

    e = pd.Series(np.nan, index=i, dtype=float)
    e = poser(e, base & present(ep3) & present(ep11), 2)
    e = poser(e, eq(n, 4) & eq(ep10c, 1) & base, 1)
    e = poser(e, eq(n, 4) & ~inlist(ep10c, 1, 2, 3) & (eq(ep38, 1) | eq(ep39, 1)) & base, 1)
    e = poser(e, inlist(n, 1, 2) & eq(inf, 1) & base, 1)
    e = poser(e, eq(n, 3) & eq(inf, 1) & eq(ep10c, 2) & base, 1)
    df["CISE_18_informel_Emp"] = e
    return df


def construire_variables(df: pd.DataFrame, journal=None) -> pd.DataFrame:
    """Enchaîne les trois do-files de construction des variables."""
    df = df.copy()
    for etape, fonction in (("1_1_Var_objectives_lower.do", variables_objectives),
                            ("1_2_Indicateur_Bulletin_To_Run.do", variables_bulletin),
                            ("Revision_CISE_12112024.do", variables_cise)):
        if journal:
            journal(f"  • {etape}")
        df = fonction(df)
    return df


# ---------------------------------------------------------------------------
# Moteurs de tabulation
# ---------------------------------------------------------------------------


def niveaux(serie: pd.Series) -> list:
    """Équivalent de ``levelsof`` : valeurs non manquantes, triées."""
    v = serie.dropna().unique().tolist()
    try:
        return sorted(v, key=float)
    except (TypeError, ValueError):
        return sorted(v, key=str)


def _somme(sub: pd.DataFrame, masque, poids: str) -> float:
    if masque is None:
        s = sub[poids]
    else:
        s = sub.loc[masque, poids]
    return float(pd.to_numeric(s, errors="coerce").sum())


def indicateur_simple(df: pd.DataFrame, varind: str, desag: list, periode: str = "trimestre",
                      effectifs: bool = False, taux: bool = False,
                      poids: str = "pmencor_ind") -> np.ndarray:
    """``export_indicateur_simple`` / ``..._annuel`` : matrice (taux | effectifs)."""
    if not effectifs and not taux:
        raise ValueError("Spécifiez au moins « effectifs » ou « taux »")
    periodes = niveaux(df[periode])
    total_lignes = sum(len(niveaux(df[d])) + 1 for d in desag) + 2
    nb = len(periodes)
    eff = np.full((total_lignes, nb), np.nan) if effectifs else None
    tx = np.full((total_lignes, nb), np.nan) if taux else None

    for j, t in enumerate(periodes):
        sub = df[df[periode] == t]
        sind = pd.to_numeric(sub[varind], errors="coerce")
        ligne = 0
        for d in desag:
            for mod in niveaux(sub[d]):
                m = sub[d] == mod
                if ligne < total_lignes:
                    if effectifs:
                        eff[ligne, j] = _somme(sub, m & (sind == 1), poids)
                    if taux:
                        denom = _somme(sub, m, poids)
                        if denom > 0:
                            tx[ligne, j] = _somme(sub, m & (sind == 1), poids) / denom * 100
                ligne += 1
            ligne += 1
        ligne += 1
        if ligne < total_lignes:
            if effectifs:
                eff[ligne, j] = _somme(sub, sind == 1, poids)
            if taux:
                denom = _somme(sub, None, poids)
                if denom > 0:
                    tx[ligne, j] = _somme(sub, sind == 1, poids) / denom * 100
    if taux and effectifs:
        return np.hstack([tx, eff])
    return tx if taux else eff


def repartition_horiz(df: pd.DataFrame, varind: str, desag: list, periode: str = "trimestre",
                      effectifs: bool = False, taux: bool = False, poids: str = "pmencor_ind",
                      lignes_corrigees: bool = False) -> np.ndarray:
    """``export_repartition_horiz`` (et la variante ANNUEL corrigée)."""
    if not effectifs and not taux:
        raise ValueError("Spécifiez au moins « effectifs » ou « taux »")
    modalites = niveaux(df[varind])
    periodes = niveaux(df[periode])
    nb_desag = len(desag)
    if lignes_corrigees:
        total_lignes = sum(len(niveaux(df[d])) for d in desag) + (nb_desag - 1) + 2
    else:
        total_lignes = sum(len(niveaux(df[d])) + 1 for d in desag) + 2
    cols_par_periode = (len(modalites) + 1) * (2 if (taux and effectifs) else 1)
    mat = np.full((total_lignes, len(periodes) * cols_par_periode), np.nan)

    def colonnes(depart: int, idx: int):
        if taux and effectifs:
            return depart + (idx - 1) * 2, depart + (idx - 1) * 2 + 1
        return (depart + idx - 1, None) if taux else (None, depart + idx - 1)

    for jp, t in enumerate(periodes):
        sub = df[df[periode] == t]
        sind = sub[varind]
        depart = jp * cols_par_periode
        for k, mod_var in enumerate(modalites, start=1):
            c_taux, c_eff = colonnes(depart, k)
            ligne = 0
            for idx_d, d in enumerate(desag, start=1):
                for mod in niveaux(sub[d]):
                    m = sub[d] == mod
                    denom = _somme(sub, m, poids)
                    numer = _somme(sub, m & (sind == mod_var), poids)
                    if ligne < total_lignes:
                        if effectifs:
                            mat[ligne, c_eff] = numer
                        if taux and denom > 0:
                            mat[ligne, c_taux] = numer / denom * 100
                    ligne += 1
                if not lignes_corrigees or idx_d < nb_desag:
                    ligne += 1
            ligne += 1
            if ligne < total_lignes:
                numer = _somme(sub, sind == mod_var, poids)
                denom = _somme(sub, None, poids)
                if effectifs:
                    mat[ligne, c_eff] = numer
                if taux and denom > 0:
                    mat[ligne, c_taux] = numer / denom * 100
        # colonne « Total » de la période
        c_taux, c_eff = colonnes(depart, len(modalites) + 1)
        ligne = 0
        for idx_d, d in enumerate(desag, start=1):
            for mod in niveaux(sub[d]):
                denom = _somme(sub, sub[d] == mod, poids)
                if ligne < total_lignes:
                    if effectifs:
                        mat[ligne, c_eff] = denom
                    if taux and denom > 0:
                        mat[ligne, c_taux] = 100
                ligne += 1
            if not lignes_corrigees or idx_d < nb_desag:
                ligne += 1
        ligne += 1
        if ligne < total_lignes:
            denom = _somme(sub, None, poids)
            if effectifs:
                mat[ligne, c_eff] = denom
            if taux and denom > 0:
                mat[ligne, c_taux] = 100
    return mat


# ---------------------------------------------------------------------------
# Écriture dans la maquette Excel
# ---------------------------------------------------------------------------


def preparer_maquette(modele: str | Path, destination: str | Path) -> Path:
    """Copie la maquette vers le fichier de résultats (``copy ... , replace``)."""
    modele, destination = Path(modele), Path(destination)
    if not modele.exists():
        raise FileNotFoundError(f"Maquette introuvable : {modele}")
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(modele, destination)
    return destination


def _coord(cellule: str) -> tuple:
    m = re.fullmatch(r"([A-Za-z]+)(\d+)", cellule.strip())
    if not m:
        raise ValueError(f"Cellule invalide : {cellule}")
    lettres, ligne = m.group(1).upper(), int(m.group(2))
    colonne = 0
    for ch in lettres:
        colonne = colonne * 26 + (ord(ch) - 64)
    return ligne, colonne


class Classeur:
    """Écrit les matrices dans la maquette, comme ``putexcel ... = matrix(RESU)``."""

    def __init__(self, chemin: str | Path):
        import openpyxl
        self.chemin = Path(chemin)
        self.wb = openpyxl.load_workbook(self.chemin)

    def feuilles(self) -> list:
        return list(self.wb.sheetnames)

    def ecrire(self, feuille: str, cellule: str, matrice: np.ndarray):
        if feuille not in self.wb.sheetnames:
            raise KeyError(f"Onglet « {feuille} » absent de la maquette")
        ws = self.wb[feuille]
        l0, c0 = _coord(cellule)
        for i in range(matrice.shape[0]):
            for j in range(matrice.shape[1]):
                v = matrice[i, j]
                if v is None or (isinstance(v, float) and np.isnan(v)):
                    continue
                try:
                    ws.cell(row=l0 + i, column=c0 + j).value = float(v)
                except (AttributeError, TypeError):
                    continue        # cellule fusionnée : ignorée

    def sauver(self):
        self.wb.save(self.chemin)
