"""Exports Excel mis en forme (charte ANSTAT) et CSV."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

VERT = "4A675A"
ORANGE = "F39422"
ORANGE_CLAIR = "FDEBD3"
ROUGE_CLAIR = "F8D7DA"
VERT_CLAIR = "DCEFE3"

_BORD = Side(style="thin", color="BFBFBF")
BORDURE = Border(left=_BORD, right=_BORD, top=_BORD, bottom=_BORD)


def nom_feuille(nom: str) -> str:
    for c in '[]:*?/\\':
        nom = nom.replace(c, "_")
    return nom[:31] or "Feuille"


def exporter_excel(chemin: str | Path, feuilles: dict[str, pd.DataFrame],
                   titre: str | None = None, surlignage: dict | None = None) -> Path:
    """Écrit un classeur avec une feuille par DataFrame, puis le met en forme.

    ``surlignage`` : {nom_feuille: fonction(ligne: pd.Series) -> "ROUGE"|"VERT"|"ORANGE"|None}
    """
    chemin = Path(chemin)
    chemin.parent.mkdir(parents=True, exist_ok=True)
    noms = {}
    with pd.ExcelWriter(chemin, engine="openpyxl") as w:
        for nom, df in feuilles.items():
            nf = nom_feuille(nom)
            d = df.copy()
            d.attrs = {}
            for c in d.columns:
                if isinstance(d[c].dtype, pd.DatetimeTZDtype):
                    d[c] = d[c].dt.tz_localize(None)
            if d.empty:
                d = pd.DataFrame({"Information": ["Aucune ligne à signaler"]}) if not len(d.columns) else d
            d.to_excel(w, sheet_name=nf, index=False, startrow=1 if titre else 0)
            noms[nf] = (df, nom)
    _mettre_en_forme(chemin, noms, titre, surlignage or {})
    return chemin


def _mettre_en_forme(chemin: Path, noms: dict, titre: str | None, surlignage: dict) -> None:
    wb = load_workbook(chemin)
    remplissage_entete = PatternFill("solid", start_color=VERT, end_color=VERT)
    couleurs = {
        "ROUGE": PatternFill("solid", start_color=ROUGE_CLAIR, end_color=ROUGE_CLAIR),
        "VERT": PatternFill("solid", start_color=VERT_CLAIR, end_color=VERT_CLAIR),
        "ORANGE": PatternFill("solid", start_color=ORANGE_CLAIR, end_color=ORANGE_CLAIR),
    }
    for nf, (df, nom) in noms.items():
        ws = wb[nf]
        ligne_entete = 2 if titre else 1
        if titre:
            ws.cell(1, 1, titre).font = Font(bold=True, size=13, color=VERT)
        for cell in ws[ligne_entete]:
            cell.fill = remplissage_entete
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BORDURE
        ws.row_dimensions[ligne_entete].height = 30
        fonction = surlignage.get(nom) or surlignage.get(nf)
        # au-delà de 5 000 lignes, on ne met pas de bordures (lenteur d'openpyxl)
        derniere = ws.max_row if (ws.max_row <= 5000 or fonction is not None) else ligne_entete
        for i, row in enumerate(ws.iter_rows(min_row=ligne_entete + 1, max_row=derniere), start=0):
            remplissage = None
            if fonction is not None and i < len(df):
                try:
                    remplissage = couleurs.get(fonction(df.iloc[i]))
                except Exception:
                    remplissage = None
            for cell in row:
                cell.border = BORDURE
                if remplissage is not None:
                    cell.fill = remplissage
        for j, col in enumerate(ws.iter_cols(min_row=ligne_entete, max_row=min(ws.max_row, 300)), start=1):
            largeur = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(j)].width = max(8, min(45, largeur + 2))
        ws.freeze_panes = ws.cell(ligne_entete + 1, 1)
        if ws.max_row > ligne_entete:
            ws.auto_filter.ref = f"A{ligne_entete}:{get_column_letter(ws.max_column)}{ws.max_row}"
    wb.save(chemin)


def exporter_csv(chemin: str | Path, df: pd.DataFrame, sep: str = ";") -> Path:
    chemin = Path(chemin)
    chemin.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(chemin, sep=sep, index=False, encoding="utf-8-sig")
    return chemin


def lire_excel_toutes_feuilles(chemin: str | Path) -> dict[str, pd.DataFrame]:
    """Relit un classeur de résultats (en ignorant une éventuelle ligne de titre)."""
    res = {}
    xls = pd.ExcelFile(chemin)
    for f in xls.sheet_names:
        df = xls.parse(f)
        # nos exports ont une ligne de titre au-dessus des en-têtes
        if df.columns.size and all(str(c).startswith("Unnamed") for c in df.columns[1:]) and len(df) > 0:
            df = xls.parse(f, header=1)
        res[f] = df
    return res
